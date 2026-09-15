import io
import pytest
import fitz  # PyMuPDF
import docx
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from fastapi.testclient import TestClient

from app.db.database import Base, get_db
from app.main import app
from app.core.security import create_access_token, get_password_hash
from app.models.user import User

# In-memory SQLite database engine for testing
SQLALCHEMY_TEST_DATABASE_URL = "sqlite:///:memory:"

test_engine = create_engine(
    SQLALCHEMY_TEST_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


@pytest.fixture(scope="function")
def db():
    """
    Fixture providing an isolated in-memory database session for each test.
    Creates all tables before test and drops them afterwards.
    """
    Base.metadata.create_all(bind=test_engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=test_engine)


@pytest.fixture(scope="function")
def client(db):
    """
    FastAPI TestClient fixture with database session override.
    """
    def _override_get_db():
        try:
            yield db
        finally:
            pass

    app.dependency_overrides[get_db] = _override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def create_test_user(db):
    """Factory fixture to create an authenticated user in the test database."""
    def _create(email: str = "testuser@example.com", password: str = "Password123!") -> dict:
        user = User(
            email=email.lower().strip(),
            password_hash=get_password_hash(password),
            is_active=True,
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        token = create_access_token(subject=user.id)
        headers = {"Authorization": f"Bearer {token}"}
        return {"user": user, "token": token, "headers": headers, "password": password}

    return _create


@pytest.fixture
def create_sample_pdf_bytes():
    """Factory fixture creating valid synthetic sample PDF resume bytes in memory."""
    def _create(text: str = None, page_count: int = 1) -> bytes:
        doc = fitz.open()
        default_text = (
            "Alex Mercer\n"
            "Email: alex.mercer@example.com | Phone: +1-555-0199 | San Francisco, CA\n"
            "PROFESSIONAL SUMMARY\n"
            "Passionate Software Engineer with 2+ years of experience building Python APIs and modern frontend web applications.\n"
            "TECHNICAL SKILLS\n"
            "Python, FastAPI, React, PostgreSQL, Docker, Git, REST APIs, Pytest\n"
            "WORK EXPERIENCE\n"
            "Software Developer - TechCorp Inc. (2022 - Present)\n"
            "- Designed scalable RESTful APIs using FastAPI and Pydantic.\n"
            "- Reduced API response latency by 40% through query optimization.\n"
            "EDUCATION\n"
            "B.S. in Computer Science - State University (2018 - 2022)\n"
            "PROJECTS\n"
            "HireLens Resume Analyzer\n"
            "- Built an automated resume parser using Python and FastAPI.\n"
            "CERTIFICATIONS\n"
            "AWS Certified Developer Associate\n"
        )
        content_text = text if text is not None else default_text

        for _ in range(page_count):
            page = doc.new_page()
            page.insert_text((50, 50), content_text)

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    return _create


@pytest.fixture
def create_sample_docx_bytes():
    """Factory fixture creating valid synthetic sample DOCX resume bytes in memory."""
    def _create(text: str = None) -> bytes:
        doc = docx.Document()
        if text:
            doc.add_paragraph(text)
        else:
            doc.add_heading("Taylor Reed", level=1)
            doc.add_paragraph("Email: taylor.reed@example.com | Phone: +1-555-0288")
            doc.add_heading("PROFILE", level=2)
            doc.add_paragraph("Entry-level Full Stack Developer seeking a challenging backend role.")
            doc.add_heading("SKILLS & TECHNOLOGIES", level=2)
            doc.add_paragraph("Python, Java, JavaScript, React, HTML, CSS, SQL")
            doc.add_heading("EXPERIENCE", level=2)
            doc.add_paragraph("Software Engineering Intern - WebSolutions (2023)")
            doc.add_paragraph("Developed interactive dashboard widgets using React.")
            doc.add_heading("EDUCATION", level=2)
            doc.add_paragraph("B.Tech in Computer Engineering - Tech Institute (2020 - 2024)")
            doc.add_heading("PROJECTS", level=2)
            doc.add_paragraph("Portfolio Tracker App - Developed full stack web app.")

            # Add sample table
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Certification"
            table.cell(0, 1).text = "Year"
            table.cell(1, 0).text = "Python Professional Cert"
            table.cell(1, 1).text = "2023"

        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()

    return _create
