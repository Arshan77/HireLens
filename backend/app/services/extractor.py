import io
import fitz  # PyMuPDF
import docx
from docx.opc.exceptions import PackageNotFoundError
import zipfile

from app.core.exceptions import FileExtractionError


def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF document using PyMuPDF (fitz).
    Handles multi-page PDFs, encrypted/password-protected PDFs, and empty PDFs.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise FileExtractionError(f"Failed to open PDF document: {str(e)}")

    if doc.is_encrypted:
        # Attempt empty password decryption
        if not doc.authenticate(""):
            doc.close()
            raise FileExtractionError("PDF document is password protected.")

    if len(doc) == 0:
        doc.close()
        raise FileExtractionError("PDF document contains no pages.")

    text_parts = []
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text:
                text_parts.append(page_text)
    except Exception as e:
        raise FileExtractionError(f"Error extracting text from PDF page: {str(e)}")
    finally:
        doc.close()

    full_text = "\n".join(text_parts).strip()
    if not full_text:
        raise FileExtractionError(
            "No extractable text found in PDF. The document may be scanned, an image, or empty."
        )

    return full_text


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a Microsoft Word (.docx) document using python-docx.
    Extracts text from both body paragraphs and table cells.
    """
    try:
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise FileExtractionError(f"Corrupted or invalid DOCX document: {str(e)}")
    except Exception as e:
        raise FileExtractionError(f"Failed to parse DOCX document: {str(e)}")

    text_parts = []

    # 1. Extract body paragraphs
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if p_text:
            text_parts.append(p_text)

    # 2. Extract table cell contents (to capture structured resume tables)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                text_parts.append(" | ".join(row_texts))

    full_text = "\n".join(text_parts).strip()
    if not full_text:
        raise FileExtractionError("No extractable text found in DOCX document.")

    return full_text


def extract_text_from_file(file_type: str, file_bytes: bytes) -> str:
    """
    Dispatch text extraction based on file type ('pdf' or 'docx').
    """
    if file_type == "pdf":
        return extract_pdf_text(file_bytes)
    elif file_type == "docx":
        return extract_docx_text(file_bytes)
    else:
        raise FileExtractionError(f"Unsupported file type '{file_type}' for extraction.")
